"""
Telegram-бот для парсинга вакансий с hh.ru
Автор: VaSeBa
Версия: 3.0 (с расширенным функционалом)
"""

# ██████████████████████████████████████████████████████████████████████
# █                        ИМПОРТ БИБЛИОТЕК                          █
# ██████████████████████████████████████████████████████████████████████

# Стандартные библиотеки
import os
import asyncio
import logging
import random
from datetime import datetime, timedelta
from typing import List, Dict, Any

# Сторонние библиотеки
import aiohttp
from dateutil import parser
from aiogram import Bot, Dispatcher, types, F
from aiogram.filters import Command
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import StatesGroup, State
from aiogram.types import FSInputFile
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo

# Локальные зависимости
from api_token import TOKEN  # Токен бота хранится в отдельном файле

# ██████████████████████████████████████████████████████████████████████
# █                         КОНФИГУРАЦИЯ                              █
# ██████████████████████████████████████████████████████████████████████

# Настройка логирования
logging.basicConfig(
    format="%(asctime)s - %(levelname)s - %(name)s - %(message)s",
    level=logging.INFO,
    handlers=[logging.StreamHandler()]
)

# Параметры API HH.ru
API_URL = "https://api.hh.ru/vacancies"
AREA = 113  # Код России
MAX_RETRIES = 3  # Попытки при ошибках сети
REQUEST_DELAY = 0.25  # Задержка между запросами (сек)
DATE_RANGE_DAYS = 30  # Период поиска (дней)
DATE_CHUNK_DAYS = 7  # Разбиение на интервалы (дней)

# Инициализация бота и диспетчера
bot = Bot(token=TOKEN)
dp = Dispatcher()

# ██████████████████████████████████████████████████████████████████████
# █                       СИСТЕМА СОСТОЯНИЙ                           █
# ██████████████████████████████████████████████████████████████████████

class ParseStates(StatesGroup):
    """Класс состояний для конечного автомата (FSM)"""
    waiting_for_profession = State()  # Ожидание ввода профессии

# ██████████████████████████████████████████████████████████████████████
# █                     ОСНОВНОЙ КЛАСС ПАРСЕРА                        █
# ██████████████████████████████████████████████████████████████████████

class HHruParser:
    """Класс для асинхронного парсинга вакансий с hh.ru"""

    def __init__(self, profession: str, chat_id: int):
        """
        Инициализация парсера
        :param profession: Искомая профессия
        :param chat_id: ID чата для отправки результатов
        """
        self.profession = profession
        self.chat_id = chat_id
        self.running = True  # Флаг для управления процессом
        self.progress_message = None  # Сообщение с прогрессом
        self.fishing_task = None  # Задача для периодических сообщений
        self.total_vacancies = 0  # Счетчик вакансий

        # Шуточные фразы для прогресс-бара
        self.fishing_phrases = [
            "🎣 Ловись рыбка - большая и маленькая!",
            "🌊 Закинули сети - ждём улова!",
            "🐠 Рыбка, иди к нам!",
            "🦈 Осторожно, акулы!",
            "🚜 Вам не нужен тракторист? А вдруг найдётся!",  # Новая шутка
            "🤖 Роботы тоже ищут работу... но пока безрезультатно",
            "📡 Сканирую секретные вакансии ЦРУ...",
            "👾 Внезапно! Вакансия для гонщика космических кораблей!"
        ]

    async def parse(self):
        """Основной метод запуска парсинга"""
        try:
            # Рассчитываем временной диапазон
            end_date = datetime.now().replace(
                hour=0, minute=0, second=0, microsecond=0
            )
            start_date = end_date - timedelta(days=DATE_RANGE_DAYS)

            # Проверка валидности дат
            if start_date >= end_date:
                await self.handle_error(ValueError("Неверный временной диапазон"))
                return

            # Генерация интервалов для поиска
            date_intervals = list(self.date_range(
                start_date,
                end_date,
                delta=timedelta(days=DATE_CHUNK_DAYS)
            ))
            total_intervals = len(date_intervals)

            if total_intervals == 0:
                await self.handle_error(ValueError("Не удалось создать интервалы поиска"))
                return

            # Запускаем фоновую задачу с шуточными сообщениями
            self.fishing_task = asyncio.create_task(self.send_fishing_phrases())

            # Отправляем начальное сообщение о прогрессе
            self.progress_message = await bot.send_message(
                self.chat_id,
                "🎣 Начинаем поиск вакансий..."
            )

            all_vacancies = []
            for i, (date_from, date_to) in enumerate(date_intervals):
                if not self.running:
                    return  # Остановка по запросу

                # Обновление прогресса
                progress = int((i + 1) / total_intervals * 100)
                await self._update_progress(progress, i, total_intervals)

                # Получение вакансий для текущего интервала
                vacancies = await self.get_vacancies(
                    date_from=date_from.isoformat(),
                    date_to=date_to.isoformat()
                )
                all_vacancies.extend(vacancies)
                await asyncio.sleep(REQUEST_DELAY)

            # Завершаем процесс
            await self.finalize_parsing(all_vacancies)

        except Exception as e:
            await self.handle_error(e)
        finally:
            # Очистка ресурсов
            if self.fishing_task:
                self.fishing_task.cancel()
            if self.progress_message:
                await self.progress_message.delete()

    async def send_fishing_phrases(self):
        """Отправка периодических шуточных сообщений"""
        while self.running:
            try:
                # Случайный выбор фразы из списка
                phrase = random.choice(self.fishing_phrases)
                await bot.send_message(self.chat_id, phrase)
                await asyncio.sleep(5)  # Интервал 5 секунд
            except asyncio.CancelledError:
                break
            except Exception as e:
                logging.error(f"Ошибка отправки сообщения: {str(e)}")
                break

    async def _update_progress(self, progress: int, current: int, total: int):
        """Обновление прогресс-бара"""
        if self.progress_message and total > 0:
            try:
                # Расчет индекса фразы для прогресса
                chunks_per_phrase = max(total // len(self.fishing_phrases), 1)
                phrase_idx = min(current // chunks_per_phrase, len(self.fishing_phrases)-1)

                # Формирование строки прогресса
                progress_bar = f"▰{'▰' * int(progress/10)}{'▱' * (10 - int(progress/10))}▰"

                # Обновление сообщения
                await self.progress_message.edit_text(
                    f"{self.fishing_phrases[phrase_idx]}\n"
                    f"{progress_bar}\n"
                    f"Прогресс: {progress}%"
                )
            except Exception as e:
                logging.error(f"Ошибка обновления прогресса: {str(e)}")

    async def get_vacancies(self, date_from: str, date_to: str) -> List[Dict]:
        """Асинхронное получение вакансий через API"""
        params = {
            "text": self.profession,
            "area": AREA,
            "date_from": date_from,
            "date_to": date_to,
            "per_page": 100,
            "page": 0
        }

        vacancies = []
        retries = MAX_RETRIES

        async with aiohttp.ClientSession() as session:
            while self.running and retries > 0:
                try:
                    async with session.get(
                        API_URL,
                        params=params,
                        timeout=aiohttp.ClientTimeout(total=10)
                    ) as response:
                        # Обработка статусных кодов
                        if response.status == 403:
                            await self.handle_rate_limit()
                            retries -= 1
                            continue

                        response.raise_for_status()
                        data = await response.json()

                        # Сбор вакансий
                        vacancies.extend(data.get("items", []))
                        if params["page"] >= data.get("pages", 1) - 1:
                            break

                        params["page"] += 1
                        await asyncio.sleep(REQUEST_DELAY)
                        retries = MAX_RETRIES

                except (aiohttp.ClientError, asyncio.TimeoutError) as e:
                    retries -= 1
                    await self.handle_network_error(e, retries)
                except Exception as e:
                    await self.handle_general_error(e)
                    break

        return vacancies

    @staticmethod
    def date_range(start: datetime, end: datetime, delta: timedelta) -> tuple:
        """Генератор интервалов дат"""
        current = start
        while current < end:
            next_date = min(current + delta, end)
            yield (current, next_date)
            current = next_date

    async def save_to_excel(self, vacancies: List[Dict]) -> str:
        """Создание Excel-файла с результатами"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Вакансии"

        # Заголовки таблицы
        headers = ["Название", "Компания", "Зарплата от", "Зарплата до",
                 "Валюта", "Регион", "Дата публикации", "Ссылка"]
        ws.append(headers)

        # Форматирование заголовков
        for col in ws[1]:
            col.font = Font(bold=True, color="FF0000")

        # Заполнение данными
        for item in vacancies:
            salary = item.get("salary") or {}
            employer = item.get("employer") or {}
            area = item.get("area") or {}

            try:
                pub_date = parser.parse(item["published_at"]).strftime("%d.%m.%Y %H:%M")
            except:
                pub_date = "N/A"

            ws.append([
                item.get("name", "Без названия"),
                employer.get("name", "Компания не указана"),
                salary.get("from", "—"),
                salary.get("to", "—"),
                salary.get("currency", "—"),
                area.get("name", "Регион не указан"),
                pub_date,
                item.get("alternate_url", "#")
            ])

        # Настройка таблицы
        tab = Table(displayName="VacanciesTable",
                   ref=f"A1:H{len(vacancies)+1}")
        tab.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium9",
            showRowStripes=True
        )
        ws.add_table(tab)

        # Автоподбор ширины столбцов
        for column in ws.columns:
            max_length = max(
                (len(str(cell.value)) for cell in column),
                default=0
            )
            adjusted_width = (max_length + 2) * 1.2
            ws.column_dimensions[column[0].column_letter].width = adjusted_width

        # Сохранение файла
        safe_name = "".join(
            [c if c.isalnum() or c in ('_', '-') else '_'
             for c in self.profession]
        ).strip('_')
        file_name = f"{safe_name}_вакансии.xlsx"
        wb.save(file_name)

        return file_name

    async def finalize_parsing(self, vacancies: List[Dict]):
        """Завершение процесса и отправка результатов"""
        if vacancies:
            try:
                file_path = await self.save_to_excel(vacancies)
                await self.send_results(file_path, len(vacancies))
            finally:
                if os.path.exists(file_path):
                    os.remove(file_path)
        else:
            await bot.send_message(self.chat_id, "😞 Не найдено подходящих вакансий")

    async def send_results(self, file_path: str, count: int):
        """Отправка файла с результатами"""
        await bot.send_document(
            self.chat_id,
            FSInputFile(file_path),
            caption=f"✅ Готово! Найдено вакансий: {count}"
        )

    # ██████████████████████████████████████████████████████████████████████
    # █                       ОБРАБОТКА ОШИБОК                           █
    # ██████████████████████████████████████████████████████████████████████

    async def handle_error(self, error: Exception):
        """Обработка критических ошибок"""
        await bot.send_message(
            self.chat_id,
            f"❌ Критическая ошибка: {str(error)}"
        )
        logging.error(f"Critical error: {str(error)}")

    async def handle_rate_limit(self):
        """Обработка превышения лимита запросов"""
        await bot.send_message(
            self.chat_id,
            "⚠️ Превышен лимит запросов! Ждем 10 секунд..."
        )
        await asyncio.sleep(10)

    async def handle_network_error(self, error: Exception, retries: int):
        """Обработка сетевых ошибок"""
        await bot.send_message(
            self.chat_id,
            f"⚠️ Ошибка сети: {str(error)}\nПопыток осталось: {retries}"
        )
        await asyncio.sleep(5)

    async def handle_general_error(self, error: Exception):
        """Обработка общих ошибок"""
        await bot.send_message(
            self.chat_id,
            f"⚠️ Ошибка: {str(error)}"
        )
        logging.error(f"Error: {str(error)}")

    def stop(self):
        """Принудительная остановка парсинга"""
        self.running = False

# ██████████████████████████████████████████████████████████████████████
# █                     ОБРАБОТЧИКИ КОМАНД                            █
# ██████████████████████████████████████████████████████████████████████

@dp.message(Command("start"))
async def cmd_start(message: types.Message):
    """Обработчик команды /start с расширенным приветствием"""
    welcome_message = (
        "👋 Приветствую, {user_name}!\n"
        "Я <b>Parser Bot</b> — твой помощник в поиске вакансий с hh.ru 🚀\n\n"
        
        "📌 <b>Основные команды:</b>\n"
        "▫️ /start — показать это сообщение\n"
        "▫️ /parse — начать поиск вакансий\n"
        "▫️ /cancel — остановить текущий поиск\n\n"
        
        "🔎 <b>Как работать с ботом:</b>\n"
        "1. Отправь команду /parse\n"
        "2. Введи название профессии (например: Сварщик)\n"
        "3. Дождись завершения поиска (~2-5 минут)\n"
        "4. Получи файл Excel с результатами 📄\n\n"
        
        "⚙️ <b>Особенности поиска:</b>\n"
        "• Ищет вакансии за последние 30 дней\n"
        "• Проверяет все регионы России\n"
        "• Поддерживает русский и английский язык\n\n"
        
        "🐛 <b>Техподдержка:</b> @vbadasin\n"
        "🐙 <b>Исходный код:</b> https://github.com/VaSeBa/HHParserTelegrammBot"
    ).format(
        user_name=message.from_user.full_name,
        days=DATE_RANGE_DAYS
    )

    await message.answer(
        text=welcome_message,
        parse_mode="HTML",
        reply_markup=types.ReplyKeyboardRemove()
    )

@dp.message(Command("parse"))
async def cmd_parse(message: types.Message, state: FSMContext):
    """Обработчик команды начала парсинга"""
    await message.answer("📝 Введите название профессии для поиска:")
    await state.set_state(ParseStates.waiting_for_profession)

@dp.message(ParseStates.waiting_for_profession)
async def process_profession(message: types.Message, state: FSMContext):
    """Обработка введенной профессии"""
    profession = message.text.strip()
    if not profession or len(profession) > 100:
        await message.answer("⚠️ Пожалуйста, введите корректное название профессии (до 100 символов)!")
        return

    await state.clear()
    parser = HHruParser(profession, message.chat.id)
    task = asyncio.create_task(parser.parse())
    await state.update_data(current_task=task)
    await message.answer(f"🔍 Начинаем поиск по запросу: {profession}")

@dp.message(Command("cancel"))
async def cmd_cancel(message: types.Message, state: FSMContext):
    """Обработчик команды отмены"""
    user_data = await state.get_data()
    if task := user_data.get('current_task'):
        task.cancel()
        await message.answer("❌ Операция отменена")
    await state.clear()

# ██████████████████████████████████████████████████████████████████████
# █                         ЗАПУСК БОТА                              █
# ██████████████████████████████████████████████████████████████████████

async def main():
    """Основная функция запуска бота"""
    await dp.start_polling(bot)

if __name__ == "__main__":
    asyncio.run(main())
